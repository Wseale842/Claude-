import csv, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
sys.path.insert(0, '.')
from lines import L

rows = [r for r in csv.reader(open('../dpgf/DPGF_lots4-8_CFA.tsv', encoding='utf-8'), delimiter='\t')]
RATES = {'mont': 40, 'tech': 60, 'ssi': 62, 'be': 75}
K_STD, K_SSI = 1.04/(1-0.20-0.08), 1.04/(1-0.15-0.08)

wb = Workbook()
thin = Side(style='thin', color='BBBBBB'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
H = PatternFill('solid', fgColor='2B2640'); HF = Font(bold=True, color='FFFFFF')
LOTF = PatternFill('solid', fgColor='E8E1FF'); PHF = PatternFill('solid', fgColor='F3F0FB')
EST = PatternFill('solid', fgColor='FFF2CC')
eur = '#,##0.00 €'

# ---- Paramètres
P = wb.active; P.title = 'Paramètres'
P.append(['Paramètre', 'Valeur', 'Commentaire'])
for r in [('Taux monteur €/h', RATES['mont'], ''), ('Taux technicien CFA €/h', RATES['tech'], ''),
          ('Taux technicien SSI €/h', RATES['ssi'], ''), ('Taux BE €/h', RATES['be'], ''),
          ('Coef k VDI/AV/vidéo/GTB', round(K_STD, 6), 'PR × 1,04 / (1 − 20 % − 8 %)'),
          ('Coef k SSI', round(K_SSI, 6), 'PR × 1,04 / (1 − 15 % − 8 %)')]:
    P.append(list(r))
for c in P[1]: c.fill, c.font = H, HF
P.column_dimensions['A'].width = 30; P.column_dimensions['C'].width = 40

# ---- Détail chiffrage
D = wb.create_sheet('Détail chiffrage')
hdr = ['Ligne DPGF', 'Lot', 'Désignation', 'U', 'Qté', 'Fourn. devis PR/u', 'Fourn. ajoutée PR/u (estimatif)',
       'h Monteur/u', 'h Tech CFA/u', 'h Tech SSI/u', 'h BE/u', 'MO PR/u', 'PR/u', 'k', 'PV/u', 'PV total', 'PR total',
       'Source fournisseur', 'Hypothèses MO / fournitures ajoutées']
D.append(hdr)
for c in D[1]: c.fill, c.font, c.alignment = H, HF, Alignment(wrap_text=True, vertical='center')
drow = {}
lot = ''
for i, r in enumerate(rows, 1):
    r += [''] * 5
    if r[0] in ('4', '5', '6', '7', '8'): lot = r[0]
    if i not in L: continue
    fd, fa, hm, ht, hs, hb, src, hyp = L[i]
    n = D.max_row + 1
    kref = "Paramètres!$B$7" if lot == '5' else "Paramètres!$B$6"
    qty = float(r[3].replace(',', '.')) if r[3] else 1
    D.append([i, lot, r[1], r[2], qty, fd, fa, hm, ht, hs, hb,
              f"=H{n}*Paramètres!$B$2+I{n}*Paramètres!$B$3+J{n}*Paramètres!$B$4+K{n}*Paramètres!$B$5",
              f"=F{n}+G{n}+L{n}", f"={kref}", f"=ROUND(M{n}*N{n},2)", f"=O{n}*E{n}", f"=M{n}*E{n}", src, hyp])
    if fa: D.cell(n, 7).fill = EST
    drow[i] = n
for col, w in zip('ABCDEFGHIJKLMNOPQRS', [7, 5, 60, 5, 6, 12, 12, 8, 8, 8, 8, 10, 11, 7, 11, 12, 12, 70, 55]):
    D.column_dimensions[col].width = w
for row in D.iter_rows(min_row=2):
    for c in row:
        c.border = B; c.alignment = Alignment(wrap_text=True, vertical='top')
    for j in (6, 7, 12, 13, 15, 16, 17): row[j-1].number_format = eur
    row[13].number_format = '0.0000'
last = D.max_row
D.freeze_panes = 'D2'

# ---- FCI (format DPGF)
F = wb.create_sheet('FCI', 0)
F.append(['KHAZNA ORY01 — Lot courants faibles base vie & chantier — FCI PROSYB v6 (DPGF FR02201-ENG-ELEC-OTH-1029-01, lots 4 à 7)'])
F.append(['Prix de vente HT. Lignes « Hors lot » / « LOT GC » non chiffrées. Cellules jaunes du détail = fournitures estimatives sans devis.'])
F.append(['N°/Lot', 'DESIGNATION', 'U', 'Qt.', 'P.U. (€)', 'P.T. (€)', 'COMMENTAIRES / SOURCE'])
for c in F[3]: c.fill, c.font = H, HF
F['A1'].font = Font(bold=True, size=12)
tot_cells = {}; start = {}
for i, r in enumerate(rows, 1):
    r += [''] * 5
    a, des, u, q = r[0], r[1], r[2], r[3]
    if a == '[1 empty row]': F.append([]); continue
    if des.startswith('TOTAL'):
        n = F.max_row + 1
        F.append(['', des, '', '', '', f"=SUM(F{start[lot]}:F{n-1})", ''])
        F.cell(n, 2).font = F.cell(n, 6).font = Font(bold=True)
        tot_cells[lot] = f"F{n}"; continue
    if a in ('4', '5', '6', '7', '8'):
        lot = a
        if lot == '8': break
        n = F.max_row + 1; start[lot] = n
        F.append([a, des]);
        for c in F[n]: c.fill = LOTF
        F.cell(n, 2).font = Font(bold=True); continue
    n = F.max_row + 1
    if i in drow:
        dr = drow[i]
        F.append([a, des, u, float(q.replace(',', '.')) if q else 1, f"='Détail chiffrage'!O{dr}", f"=D{n}*E{n}",
                  f"='Détail chiffrage'!R{dr}"])
    else:
        F.append([a, des, u, q, '', '', 'Hors PROSYB' if a in ('Hors lot', 'LOT GC') else ''])
        if a.startswith('4.') or a.startswith('5.') or a.startswith('6.') or (not u and not q):
            for c in F[n]: c.fill = PHF
            F.cell(n, 2).font = Font(bold=True, italic=True)
n = F.max_row + 2
F.cell(n, 2, 'TOTAL GÉNÉRAL LOTS 4 À 7 HT').font = Font(bold=True)
F.cell(n, 6, '=' + '+'.join(tot_cells[k] for k in ('4', '5', '6', '7'))).font = Font(bold=True)
grand = f"F{n}"
for row in F.iter_rows(min_row=3):
    for c in row: c.border = B; c.alignment = Alignment(wrap_text=True, vertical='top')
    row[4].number_format = row[5].number_format = eur
for col, w in zip('ABCDEFG', [9, 75, 6, 7, 13, 14, 80]): F.column_dimensions[col].width = w
F.freeze_panes = 'A4'

# ---- Récap
R = wb.create_sheet('Récap', 0)
R.append(['Récap FCI v6 — KHAZNA ORY01 — courants faibles base vie & chantier — 16/10/2026 (mise à jour GTB AETHERYS)'])
R['A1'].font = Font(bold=True, size=12)
R.append([])
R.append(['Lot', 'Désignation', 'Fournitures PR', 'MO PR', 'Total PR', 'Total PV HT', 'Marge brute', 'Taux / PV'])
for c in R[3]: c.fill, c.font = H, HF
names = {'4': 'Informatique, audiovisuel, affichage', '5': 'Sécurité SSI', '6': 'Vidéosurveillance', '7': 'GTB / GTE (offre AETHERYS 16/10)'}
D_ = "'Détail chiffrage'"
for k in ('4', '5', '6', '7'):
    n = R.max_row + 1
    R.append([int(k), names[k],
              f"=SUMPRODUCT(({D_}!B2:B{last}=\"{k}\")*({D_}!F2:F{last}+{D_}!G2:G{last})*{D_}!E2:E{last})",
              f"=SUMPRODUCT(({D_}!B2:B{last}=\"{k}\")*{D_}!L2:L{last}*{D_}!E2:E{last})",
              f"=C{n}+D{n}", f"=FCI!{tot_cells[k]}", f"=F{n}-E{n}", f"=IF(F{n}=0,0,G{n}/F{n})"])
n = R.max_row + 1
R.append(['', 'TOTAL', f"=SUM(C4:C7)", f"=SUM(D4:D7)", f"=SUM(E4:E7)", f"=SUM(F4:F7)", f"=SUM(G4:G7)", f"=G{n}/F{n}"])
for c in R[n]: c.font = Font(bold=True)
for row in R.iter_rows(min_row=4, max_row=n):
    for j in range(2, 7): row[j].number_format = eur
    row[7].number_format = '0.0%'
    for c in row: c.border = B
R.append([])
R.append(['Hors FCI : études EXE.CFA (DPGF §2 « Architecture IT & audiovisuel »), part CFA de la maintenance 28 mois et permanence 24/7 (DPGF §8), encadrement, déplacements, nacelle.'])
for col, w in zip('ABCDEFGH', [6, 40, 15, 15, 15, 15, 15, 10]): R.column_dimensions[col].width = w

# ---- Points ouverts
O = wb.create_sheet('Points ouverts')
O.append(['#', 'Sujet', 'Action', 'Impact prix'])
for c in O[1]: c.fill, c.font = H, HF
for p in [
 ('GTB', 'TMA, licence CARE et abonnement 5G AETHERYS sur 24 mois alors que chantier 29 mois (maintenance DPGF 28 mois)', 'Demander à AETHERYS +5 mois ; provision 870 € PR intégrée', '+'),
 ('GTB', 'Connecteur RS485 750-960 cité dans le mémoire AETHERYS mais absent de son DPGF ; 3e bloc libellé « TGBT HUB 2 » au lieu de HUB 3 et « coupleur de bus » au lieu de PFC100', 'Faire corriger l\'offre AETHERYS', 'faible'),
 ('GTB', 'Alertes SMS : non décrites par AETHERYS (Ignition Edge = e-mail natif) alors que DPGF demande SMS/mail', 'Confirmer passerelle SMS (IXON ou module Ignition) ou retirer SMS', '+250 € si passerelle'),
 ('GTB', 'Routeur IXON 5G = sortie internet du réseau terrain : contradiction avec « aucun accès internet direct » du mémoire', 'Placer en VLAN 910 VPN de maintenance, activation sur validation EXYTE', '0'),
 ('Vidéo', 'Caméras DS-2CD1643G2-LIZSU : fiche = détection de mouvement humain/véhicule seulement (pas de franchissement de ligne/intrusion de zone listés) ; NVR : protection périmétrique limitée à 2 voies 4 MP ; port des EPI non supporté', 'Demander à CCF variante AcuSense (ex. gamme 2xx6G2-I…SU/SL) ou reformuler l\'offre', '+ à chiffrer'),
 ('Vidéo', 'NVR DS-7732NXI-K4 : 16 To max par disque selon fiche ; disque 20 To proposé + disque unique = point unique de panne', 'Remplacer par 2 × 10 To (ou 4 × 8 To) compatibles', '≈ neutre'),
 ('AV', 'Micro LOG025 MeetUp incompatible Rally Bar → 4 Rally Mic Pod 989-000430 ; chariot Chief 50-85" vs écran 86" ; totems ERARD prévus OH55F (portrait) vs OHDX', 'Recotation Sonepar attendue', '±'),
 ('AV', 'Serveur Navori QL et durée licence Navori non chiffrés', 'Prix Sonepar', '1 100 € estimatif'),
 ('VDI', 'Pose cuivre et certification classe EA non incluses Octet Plus ; coffrets IP55 retirés (SPIE) → devis révisé à demander', 'Devis révisé Octet Plus', '0'),
 ('SSI', 'Câbles CR1-C1, coffret de reports : prix de repli estimatifs', 'Consultation L19 Sonepar en cours', '±'),
 ('Validité', 'Devis fournisseurs valables 27-31/10/2026 alors que Haskoning exige ≥ 90 jours de validité', 'Demander prolongation ou clause de révision', 'risque'),
]:
    O.append([O.max_row] + list(p))
for col, w in zip('ABCDE', [4, 10, 90, 55, 18]): O.column_dimensions[col].width = w
for row in O.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = B

wb.save('FCI_Khazna_CFA_SSI_AV_IT_v6_AETHERYS.xlsx')
print('saved', last, 'detail rows')
